"""
MCP Second Brain Server for Home Assistant v3.0.0 (MCP Python SDK 2.x)

Mehrere voneinander getrennte "Brains" (Standard: privat, beruf) unter einem
Stammordner (Standard /share/second_brain). Jedes Brain folgt dem LLM-Wiki-Muster:

    <brain>/CLAUDE.md          Regeln (Schema)
    <brain>/raw/inbox/         Eingangskorb für neue Quellen
    <brain>/raw/...            Originale (unveränderlich)
    <brain>/wiki/start.md      Cockpit
    <brain>/wiki/index.md      Katalog aller Seiten
    <brain>/wiki/log.md        Chronik (nur anhängen)
    <brain>/wiki/{personen,projekte,themen,konzepte,quellen}/

Jeder Tool-Aufruf nennt sein Brain; nichts ist brainübergreifend erreichbar.
Gelöschtes wandert in <brain>/.trash.
"""

import base64
import io
import json
import os
import re
import shutil
import zipfile
from datetime import datetime, date
from pathlib import Path

from mcp.server.mcpserver import MCPServer
from mcp.server.mcpserver.exceptions import ToolError
from mcp.server.transport_security import TransportSecuritySettings
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.requests import Request
from starlette.responses import Response
import uvicorn

VERSION = "3.0.0"

# ── Konfiguration ─────────────────────────────────────────────────────────────

TOKEN = os.environ.get("MCP_TOKEN", "")
BASE = Path(os.environ.get("SB_ROOT", "/share/second_brain")).resolve()
PORT = int(os.environ.get("SB_PORT", "8773"))
READ_ONLY = os.environ.get("SB_READ_ONLY", "false").lower() == "true"
MAX_FILE_BYTES = int(os.environ.get("SB_MAX_FILE_KB", "512")) * 1024
MAX_UPLOAD_BYTES = int(os.environ.get("SB_MAX_UPLOAD_MB", "50")) * 1024 * 1024
BRAINS = [b.strip() for b in os.environ.get("SB_BRAINS", "privat,beruf").split(",") if b.strip()]
# Optionale externe Eingangsordner (z. B. OneDrive-Sync), je Brain ein Unterordner: <dir>/<brain>
IMPORT_DIR = os.environ.get("SB_IMPORT_DIR", "").strip()

TEXT_EXT = {".md", ".txt", ".json", ".yaml", ".yml", ".csv"}
EXTRACT_EXT = {".pdf", ".docx", ".xlsx", ".html", ".htm"}
TRASH_DIR = ".trash"
RULES = "CLAUDE.md"
WIKI_DIRS = ["personen", "projekte", "themen", "konzepte", "quellen"]
LINK_RE = re.compile(r"\[\[([^\]|#]+)(?:#[^\]|]*)?(?:\|[^\]]*)?\]\]")
STAND_RE = re.compile(r"^stand:\s*(\d{4}-\d{2}-\d{2})", re.M)

BASE.mkdir(parents=True, exist_ok=True)


def _ensure_structure(brain: str) -> None:
    root = BASE / brain
    for d in ["raw/inbox"] + [f"wiki/{w}" for w in WIKI_DIRS]:
        (root / d).mkdir(parents=True, exist_ok=True)


if not READ_ONLY:
    for _b in BRAINS:
        _ensure_structure(_b)

print(f"[Second Brain] v{VERSION} base={BASE} brains={BRAINS} read_only={READ_ONLY} "
      f"import_dir={IMPORT_DIR or '-'} token={'enabled' if TOKEN else 'DISABLED'}", flush=True)

# ── Server ────────────────────────────────────────────────────────────────────

INSTRUCTIONS = (
    "Second Brain: Jonathans persönliches Wissens-Wiki nach dem LLM-Wiki-Muster (Karpathy). "
    f"Es gibt getrennte Brains: {', '.join(BRAINS)}. Jeder Aufruf braucht den Parameter brain; "
    "Inhalte werden NIE zwischen Brains kopiert oder verknüpft. Wahl: Berufliches (SPS, TIA, Ceracon, "
    "Kollegen) -> beruf, alles andere -> privat. Im Brain beruf nie Firmeninterna aus Notion/Outlook/Teams "
    "ablegen, nur Wissen aus den Chats. "
    "ABLAUF: 1) Zu Beginn jeder Aufgabe, die Jonathans Projekte, Geräte, Entscheidungen, Ziele oder "
    "Vorlieben betrifft, ZUERST sb_start(brain) aufrufen: liefert Regeln (CLAUDE.md), Cockpit, Index, "
    "Eingangskorb und letzte Log-Einträge. Die Regeln in CLAUDE.md sind verbindlich. "
    "2) Liegt etwas im Eingangskorb, das zuerst verarbeiten (Ingest). "
    "3) Vor dem Ende jeder Unterhaltung mit neuen Fakten, Entscheidungen, Plänen oder Vorlieben die "
    "betroffenen Wiki-Seiten aktualisieren und sb_log aufrufen - ohne dass Jonathan daran erinnert. "
    "Mehrere Seiten auf einmal: sb_read_many / sb_write_many. Keine Passwörter, Tokens oder Zugangsdaten speichern."
)

mcp = MCPServer(name="Second Brain", instructions=INSTRUCTIONS, version=VERSION)

# ── Hilfsfunktionen ───────────────────────────────────────────────────────────


def _root(brain: str) -> Path:
    b = (brain or "").strip().lower()
    if b not in BRAINS:
        raise ToolError(f"Unbekanntes Brain '{brain}'. Erlaubt: {BRAINS}")
    return (BASE / b).resolve()


def _rel(root: Path, p: Path) -> str:
    return p.relative_to(root).as_posix() or "."


def _resolve(root: Path, path: str) -> Path:
    path = (path or "").strip().replace("\\", "/")
    if path.startswith(str(root)):
        path = path[len(str(root)):]
    path = path.lstrip("/")
    resolved = (root / path).resolve()
    if resolved != root and root not in resolved.parents:
        raise ToolError(f"Pfad '{path}' liegt außerhalb dieses Brains.")
    return resolved


def _in_trash(root: Path, p: Path) -> bool:
    return TRASH_DIR in p.relative_to(root).parts


def _is_raw(root: Path, p: Path) -> bool:
    parts = p.relative_to(root).parts
    return bool(parts) and parts[0] == "raw"


def _check_text_file(root: Path, path: str) -> Path:
    p = _resolve(root, path)
    if p == root:
        raise ToolError("Pfad zeigt auf den Stammordner, nicht auf eine Datei.")
    if p.suffix.lower() not in TEXT_EXT:
        raise ToolError(f"Dateityp '{p.suffix}' hier nicht erlaubt. Text: {sorted(TEXT_EXT)}")
    if _in_trash(root, p):
        raise ToolError("Im Papierkorb wird nicht geschrieben.")
    return p


def _check_any_file(root: Path, path: str) -> Path:
    """Beliebiger Dateityp - aber nur unter raw/ (für Originale). Text überall."""
    p = _resolve(root, path)
    if p == root:
        raise ToolError("Pfad zeigt auf den Stammordner, nicht auf eine Datei.")
    if p.suffix.lower() not in TEXT_EXT and not _is_raw(root, p):
        raise ToolError("Nicht-Text-Dateien sind nur unter raw/ erlaubt.")
    return p


def _check_writable() -> None:
    if READ_ONLY:
        raise ToolError("Second Brain ist im Nur-Lese-Modus (Option read_only).")


def _check_size(content: str) -> None:
    if len(content.encode("utf-8")) > MAX_FILE_BYTES:
        raise ToolError(f"Inhalt größer als {MAX_FILE_BYTES // 1024} KB - bitte aufteilen.")


def _mtime(p: Path) -> str:
    return datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")


def _iter_files(root: Path, base: Path, exts=TEXT_EXT, include_trash=False):
    for f in sorted(base.rglob("*")):
        if not f.is_file() or (exts is not None and f.suffix.lower() not in exts):
            continue
        if not include_trash and _in_trash(root, f):
            continue
        if any(part.startswith(".") for part in f.relative_to(root).parts):
            continue
        yield f


def _read_text(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="replace")


def _extract(p: Path) -> str:
    """Text aus PDF/DOCX/XLSX/HTML ziehen."""
    ext = p.suffix.lower()
    if ext in TEXT_EXT:
        return _read_text(p)
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(p))
        return "\n\n".join(f"--- Seite {i + 1} ---\n{(pg.extract_text() or '').strip()}"
                           for i, pg in enumerate(r.pages))
    if ext == ".docx":
        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", "\t", xml)
        return re.sub(r"<[^>]+>", "", xml)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(p), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"--- Blatt: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    out.append("\t".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    if ext in (".html", ".htm"):
        t = _read_text(p)
        t = re.sub(r"(?is)<(script|style).*?</\1>", "", t)
        t = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", t)
        return re.sub(r"<[^>]+>", "", t)
    raise ToolError(f"Kein Textauszug für '{ext}' möglich.")


def _frontmatter(text: str) -> dict:
    if not text.startswith("---"):
        return {}
    end = text.find("\n---", 3)
    if end < 0:
        return {}
    fm = {}
    for line in text[3:end].splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            fm[k.strip().lower()] = v.strip()
    return fm


def _wiki_pages(root: Path) -> dict:
    """name -> Pfad; Schlüssel: 'projekte/x' und 'x'."""
    wiki = root / "wiki"
    pages = {}
    for f in _iter_files(root, wiki, exts={".md"}):
        rel = f.relative_to(wiki).with_suffix("").as_posix()
        pages[rel.lower()] = f
        pages.setdefault(f.stem.lower(), f)
    return pages


def _links(text: str):
    return [m.group(1).strip().lower().removesuffix(".md") for m in LINK_RE.finditer(text)]


def _slice(text: str, offset: int, limit: int) -> dict:
    total = len(text)
    offset = max(0, offset)
    chunk = text[offset:offset + limit] if limit > 0 else text[offset:]
    return {"content": chunk, "total_chars": total, "offset": offset,
            "next_offset": offset + len(chunk) if offset + len(chunk) < total else None}

# ── Tools: Einstieg ───────────────────────────────────────────────────────────


@mcp.tool()
def sb_start(brain: str) -> dict:
    """ZUERST aufrufen. Liefert in EINEM Aufruf: Regeln (CLAUDE.md), Cockpit (wiki/start.md),
    Index (wiki/index.md), Dateien im Eingangskorb (raw/inbox + externer Eingang),
    die letzten Log-Einträge und Seitenzahlen. Args: brain ('privat' | 'beruf')."""
    root = _root(brain)

    def rd(rel):
        p = root / rel
        return _read_text(p) if p.is_file() else None

    inbox = [{"path": _rel(root, f), "size": f.stat().st_size, "modified": _mtime(f)}
             for f in _iter_files(root, root / "raw" / "inbox", exts=None)]
    log = rd("wiki/log.md") or ""
    entries = re.split(r"(?m)^(?=## \[)", log)
    recent_log = "".join(entries[-8:]).strip() if len(entries) > 1 else ""
    pages = sum(1 for _ in _iter_files(root, root / "wiki", exts={".md"}))
    return {
        "brain": brain, "version": VERSION, "read_only": READ_ONLY, "today": date.today().isoformat(),
        "rules": rd(RULES), "cockpit": rd("wiki/start.md"), "index": rd("wiki/index.md"),
        "inbox": inbox, "external_inbox": _external_inbox(brain), "recent_log": recent_log,
        "wiki_pages": pages,
    }

# ── Tools: Lesen ──────────────────────────────────────────────────────────────


@mcp.tool()
def sb_read(brain: str, path: str, offset: int = 0, limit: int = 60000) -> dict:
    """Liest eine Datei. Text direkt; PDF/DOCX/XLSX/HTML unter raw/ als extrahierter Text.
    Große Inhalte seitenweise über offset/limit (Zeichen); next_offset zeigt, ob mehr kommt.
    Args: brain, path (relativ, z. B. 'wiki/projekte/hundebox.md'), offset, limit."""
    root = _root(brain)
    p = _resolve(root, path)
    if not p.is_file():
        return {"success": False, "error": f"Datei existiert nicht: {path}"}
    ext = p.suffix.lower()
    if ext not in TEXT_EXT and ext not in EXTRACT_EXT:
        return {"success": False, "path": _rel(root, p), "size": p.stat().st_size,
                "error": f"Binärdatei '{ext}' ohne Textauszug (Original bleibt unter raw/)."}
    text = _extract(p)
    return {"success": True, "path": _rel(root, p), "modified": _mtime(p), **_slice(text, offset, limit)}


@mcp.tool()
def sb_read_many(brain: str, paths: list[str], max_chars_each: int = 20000) -> dict:
    """Liest mehrere Dateien in einem Aufruf (schneller als viele sb_read).
    Args: brain, paths (Liste), max_chars_each (Kürzung je Datei)."""
    root = _root(brain)
    out = []
    for path in paths[:40]:
        try:
            p = _resolve(root, path)
            if not p.is_file():
                out.append({"path": path, "error": "existiert nicht"})
                continue
            text = _extract(p)
            out.append({"path": _rel(root, p), "modified": _mtime(p),
                        "content": text[:max_chars_each], "truncated": len(text) > max_chars_each})
        except Exception as e:  # noqa: BLE001
            out.append({"path": path, "error": str(e)})
    return {"success": True, "files": out}


@mcp.tool()
def sb_list(brain: str, folder: str = "", recursive: bool = False) -> dict:
    """Listet Dateien und Unterordner eines Brains. Args: brain, folder (relativ), recursive."""
    root = _root(brain)
    base = _resolve(root, folder)
    if not base.is_dir():
        return {"success": False, "error": f"Ordner existiert nicht: {folder}"}
    entries = []
    items = base.rglob("*") if recursive else base.iterdir()
    for item in sorted(items):
        if _in_trash(root, item) or any(x.startswith(".") for x in item.relative_to(root).parts):
            continue
        if item.is_dir():
            entries.append({"path": _rel(root, item), "type": "folder"})
        else:
            entries.append({"path": _rel(root, item), "type": "file",
                            "size": item.stat().st_size, "modified": _mtime(item)})
    return {"success": True, "folder": _rel(root, base), "entries": entries}


@mcp.tool()
def sb_search(brain: str, query: str, folder: str = "wiki", max_results: int = 20) -> dict:
    """Volltextsuche (Groß/Klein egal) über Dateinamen und Inhalte der Textdateien.
    Alle Wörter müssen vorkommen. Treffer im Titel/Dateinamen zählen mehr.
    Args: brain, query, folder (Standard 'wiki'; '' = ganzes Brain inkl. raw), max_results."""
    root = _root(brain)
    terms = [t.lower() for t in query.split() if t.strip()]
    if not terms:
        return {"success": False, "error": "Leere Suche."}
    base = _resolve(root, folder)
    results = []
    for f in _iter_files(root, base):
        text = _read_text(f)
        low = text.lower()
        rel = _rel(root, f)
        hay = rel.lower() + "\n" + low
        if not all(t in hay for t in terms):
            continue
        lines = text.splitlines()
        hits = [{"line": i + 1, "text": ln.strip()[:240]}
                for i, ln in enumerate(lines) if any(t in ln.lower() for t in terms)][:5]
        title = next((ln for ln in lines if ln.startswith("# ")), "").lower()
        score = (sum(min(low.count(t), 20) for t in terms)
                 + 8 * sum(t in rel.lower() for t in terms) + 6 * sum(t in title for t in terms))
        results.append({"path": rel, "modified": _mtime(f), "score": score, "hits": hits})
    results.sort(key=lambda r: r["score"], reverse=True)
    return {"success": True, "query": query, "total": len(results), "results": results[:max(1, max_results)]}


@mcp.tool()
def sb_backlinks(brain: str, page: str) -> dict:
    """Welche Wiki-Seiten verlinken auf diese Seite? Args: brain, page (z. B. 'projekte/hundebox')."""
    root = _root(brain)
    target = page.strip().lower().removesuffix(".md").removeprefix("wiki/")
    stem = target.split("/")[-1]
    refs = []
    for f in _iter_files(root, root / "wiki", exts={".md"}):
        ls = _links(_read_text(f))
        if target in ls or stem in ls:
            refs.append(_rel(root, f))
    return {"success": True, "page": target, "backlinks": refs}


@mcp.tool()
def sb_recent(brain: str, limit: int = 15) -> dict:
    """Zuletzt geänderte Dateien (neueste zuerst). Args: brain, limit."""
    root = _root(brain)
    files = sorted(_iter_files(root, root), key=lambda f: f.stat().st_mtime, reverse=True)
    return {"success": True, "files": [{"path": _rel(root, f), "modified": _mtime(f)} for f in files[:max(1, limit)]]}

# ── Tools: Schreiben ──────────────────────────────────────────────────────────


def _write(root: Path, path: str, content: str, overwrite: bool) -> dict:
    p = _check_text_file(root, path)
    if _is_raw(root, p) and p.exists():
        raise ToolError("Originale unter raw/ werden nie überschrieben.")
    _check_size(content)
    if p.exists() and not overwrite:
        return {"success": False, "path": _rel(root, p),
                "error": "Datei existiert bereits. sb_append/sb_edit nutzen oder overwrite=true."}
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return {"success": True, "path": _rel(root, p), "bytes": len(content.encode("utf-8"))}


@mcp.tool()
def sb_write(brain: str, path: str, content: str, overwrite: bool = False) -> dict:
    """Legt eine Textdatei an (Ordner werden erzeugt). Bestehende nur mit overwrite=true.
    Args: brain, path, content, overwrite."""
    _check_writable()
    return _write(_root(brain), path, content, overwrite)


@mcp.tool()
def sb_write_many(brain: str, files: list[dict]) -> dict:
    """Schreibt mehrere Dateien in einem Aufruf (z. B. beim Ingest 10-15 Seiten).
    files: [{"path": "...", "content": "...", "overwrite": true/false}, ...] (max. 40). Args: brain, files."""
    _check_writable()
    root = _root(brain)
    results = []
    for f in files[:40]:
        try:
            results.append(_write(root, f.get("path", ""), f.get("content", ""), bool(f.get("overwrite", False))))
        except ToolError as e:
            results.append({"success": False, "path": f.get("path"), "error": str(e)})
    return {"success": all(r.get("success") for r in results), "results": results}


@mcp.tool()
def sb_append(brain: str, path: str, content: str, heading: str = "") -> dict:
    """Hängt Text an (legt die Datei an, falls nötig). Mit heading (z. B. '## Verlauf') wird am Ende
    dieses Abschnitts eingefügt; fehlt er, wird er am Dateiende angelegt. Args: brain, path, content, heading."""
    _check_writable()
    root = _root(brain)
    p = _check_text_file(root, path)
    old = _read_text(p) if p.exists() else ""
    block = content.rstrip("\n") + "\n"
    if heading.strip():
        h = heading.strip()
        level = len(h) - len(h.lstrip("#")) or 2
        if not h.startswith("#"):
            h = "#" * level + " " + h
        lines = old.splitlines(keepends=True)
        idx = next((i for i, ln in enumerate(lines) if ln.strip() == h), None)
        if idx is None:
            new = old.rstrip("\n") + ("\n\n" if old.strip() else "") + h + "\n" + block
        else:
            end = len(lines)
            for j in range(idx + 1, len(lines)):
                m = re.match(r"^(#+)\s", lines[j])
                if m and len(m.group(1)) <= level:
                    end = j
                    break
            before = "".join(lines[:end]).rstrip("\n") + "\n"
            after = "".join(lines[end:])
            new = before + block + ("\n" + after if after else "")
    else:
        new = (old.rstrip("\n") + "\n" if old.strip() else "") + block
    _check_size(new)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(new, encoding="utf-8")
    return {"success": True, "path": _rel(root, p), "created": not old}


@mcp.tool()
def sb_edit(brain: str, path: str, old_text: str, new_text: str) -> dict:
    """Ersetzt eine genau einmal vorkommende Textstelle (leerer new_text löscht sie).
    Args: brain, path, old_text, new_text."""
    _check_writable()
    root = _root(brain)
    p = _check_text_file(root, path)
    if not p.is_file():
        return {"success": False, "error": f"Datei existiert nicht: {path}"}
    if _is_raw(root, p) and not p.relative_to(root).parts[1:2] == ("inbox",):
        raise ToolError("Originale unter raw/ werden nicht verändert.")
    text = _read_text(p)
    n = text.count(old_text) if old_text else 0
    if n != 1:
        return {"success": False, "error": f"old_text kommt {n}-mal vor (muss genau 1-mal sein)."}
    new = text.replace(old_text, new_text, 1)
    _check_size(new)
    p.write_text(new, encoding="utf-8")
    return {"success": True, "path": _rel(root, p)}


@mcp.tool()
def sb_log(brain: str, kind: str, text: str) -> dict:
    """Hängt einen Eintrag an wiki/log.md an: '## [JJJJ-MM-TT HH:MM] kind | text'.
    kind: ingest | query | update | lint | setup. Args: brain, kind, text."""
    _check_writable()
    root = _root(brain)
    p = root / "wiki" / "log.md"
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    first, *rest = text.strip().splitlines() or [""]
    entry = f"## [{stamp}] {kind.strip().lower()} | {first}\n" + ("\n".join(rest) + "\n" if rest else "")
    old = _read_text(p) if p.exists() else "# Log\n"
    new = old.rstrip("\n") + "\n\n" + entry
    _check_size(new)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(new, encoding="utf-8")
    return {"success": True, "entry": entry.strip()}


@mcp.tool()
def sb_move(brain: str, source: str, destination: str) -> dict:
    """Verschiebt/benennt um (überschreibt nie). Nicht-Text-Dateien nur innerhalb raw/
    (z. B. raw/inbox/x.pdf -> raw/2026/x.pdf nach dem Ingest). Args: brain, source, destination."""
    _check_writable()
    root = _root(brain)
    src = _check_any_file(root, source)
    dst = _check_any_file(root, destination)
    if not src.is_file():
        return {"success": False, "error": f"Quelle existiert nicht: {source}"}
    if dst.exists():
        return {"success": False, "error": f"Ziel existiert bereits: {destination}"}
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), str(dst))
    return {"success": True, "source": _rel(root, src), "destination": _rel(root, dst)}


@mcp.tool()
def sb_delete(brain: str, path: str) -> dict:
    """Verschiebt eine Datei in den Papierkorb (<brain>/.trash/<Zeitstempel>/...). Nichts wird
    endgültig gelöscht. Args: brain, path."""
    _check_writable()
    root = _root(brain)
    p = _check_any_file(root, path)
    if not p.is_file():
        return {"success": False, "error": f"Datei existiert nicht: {path}"}
    if _in_trash(root, p):
        return {"success": False, "error": "Datei liegt bereits im Papierkorb."}
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dst = root / TRASH_DIR / stamp / p.relative_to(root)
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(p), str(dst))
    return {"success": True, "path": _rel(root, p), "trashed_to": _rel(root, dst)}


@mcp.tool()
def sb_upload(brain: str, path: str, data_base64: str, append: bool = False) -> dict:
    """Legt eine Originaldatei (auch binär, z. B. PDF) unter raw/ ab, typischerweise raw/inbox/.
    Große Dateien in Stücken: erstes Stück append=false, weitere append=true.
    Args: brain, path (muss mit 'raw/' beginnen), data_base64, append."""
    _check_writable()
    root = _root(brain)
    p = _resolve(root, path)
    if not _is_raw(root, p) or p == root / "raw":
        raise ToolError("Uploads nur unter raw/ (z. B. raw/inbox/datei.pdf).")
    try:
        data = base64.b64decode(data_base64, validate=True)
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Ungültiges Base64: {e}")
    if p.exists() and not append:
        return {"success": False, "error": "Datei existiert bereits (Originale werden nie überschrieben)."}
    size = (p.stat().st_size if p.exists() else 0) + len(data)
    if size > MAX_UPLOAD_BYTES:
        raise ToolError(f"Datei größer als {MAX_UPLOAD_BYTES // 1024 // 1024} MB.")
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "ab" if append else "wb") as fh:
        fh.write(data)
    return {"success": True, "path": _rel(root, p), "size": size}

# ── Externer Eingang (optional, z. B. OneDrive-Sync) ──────────────────────────


def _import_state(brain: str) -> tuple[Path, dict]:
    sp = _root(brain) / ".imported.json"
    try:
        return sp, json.loads(sp.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return sp, {}


def _external_inbox(brain: str) -> list:
    if not IMPORT_DIR:
        return []
    d = Path(IMPORT_DIR) / brain
    if not d.is_dir():
        return []
    _, state = _import_state(brain)
    out = []
    for f in sorted(d.rglob("*")):
        if f.is_file() and not f.name.startswith("."):
            key = f.relative_to(d).as_posix()
            sig = f"{f.stat().st_size}:{int(f.stat().st_mtime)}"
            if state.get(key) != sig:
                out.append({"name": key, "size": f.stat().st_size, "modified": _mtime(f)})
    return out


@mcp.tool()
def sb_import(brain: str, name: str) -> dict:
    """Kopiert eine neue Datei aus dem externen Eingangsordner (Option import_dir/<brain>) nach
    raw/inbox/. Das Original im externen Ordner bleibt unangetastet und wird als importiert gemerkt.
    Args: brain, name (wie in sb_start.external_inbox angezeigt)."""
    _check_writable()
    if not IMPORT_DIR:
        raise ToolError("Kein externer Eingangsordner konfiguriert (Option import_dir).")
    root = _root(brain)
    d = (Path(IMPORT_DIR) / brain).resolve()
    src = (d / name).resolve()
    if d not in src.parents or not src.is_file():
        raise ToolError(f"Datei nicht im externen Eingang: {name}")
    dst = root / "raw" / "inbox" / src.name
    if dst.exists():
        dst = dst.with_name(f"{dst.stem}_{datetime.now():%Y%m%d%H%M%S}{dst.suffix}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    sp, state = _import_state(brain)
    state[src.relative_to(d).as_posix()] = f"{src.stat().st_size}:{int(src.stat().st_mtime)}"
    sp.write_text(json.dumps(state, indent=1), encoding="utf-8")
    return {"success": True, "copied_to": _rel(root, dst)}

# ── Aufräumen (Lint) ──────────────────────────────────────────────────────────


@mcp.tool()
def sb_lint(brain: str, stale_days: int = 180) -> dict:
    """Mechanische Gesundheitsprüfung des Wikis: kaputte Links, verwaiste Seiten (ohne eingehende
    Links), Seiten fehlen im Index, Seiten ohne Quellen, veraltetes 'stand:', offene Fragen,
    markierte Widersprüche, volle Eingangskörbe. Inhaltliche Prüfung macht Claude danach.
    Args: brain, stale_days."""
    root = _root(brain)
    wiki = root / "wiki"
    pages = _wiki_pages(root)
    special = {"start", "index", "log"}
    index_text = (_read_text(wiki / "index.md") if (wiki / "index.md").exists() else "").lower()
    index_links = set(_links(index_text))
    inbound = {}
    broken, no_source, stale, questions, conflicts, not_indexed = [], [], [], [], [], []
    today = date.today()
    files = list(_iter_files(root, wiki, exts={".md"}))
    for f in files:
        rel = f.relative_to(wiki).with_suffix("").as_posix().lower()
        text = _read_text(f)
        for ln in set(_links(text)):
            target = pages.get(ln)
            if target is None:
                broken.append({"page": rel, "link": ln})
            elif target != f:
                inbound.setdefault(target, set()).add(rel)
        if f.stem.lower() in special:
            continue
        fm = _frontmatter(text)
        if not fm.get("quellen") and "## quellen" not in text.lower():
            no_source.append(rel)
        m = STAND_RE.search(text)
        if m:
            try:
                age = (today - date.fromisoformat(m.group(1))).days
                if age > stale_days:
                    stale.append({"page": rel, "stand": m.group(1), "days": age})
            except ValueError:
                pass
        q = sum(1 for ln in text.splitlines() if re.search(r"❓|offene frage", ln, re.I))
        if q:
            questions.append({"page": rel, "count": q})
        c = sum(1 for ln in text.splitlines() if re.search(r"⚠️|widerspruch", ln, re.I))
        if c:
            conflicts.append({"page": rel, "count": c})
        if rel not in index_links and f.stem.lower() not in index_links:
            not_indexed.append(rel)
    orphans = [f.relative_to(wiki).with_suffix("").as_posix() for f in files
               if f.stem.lower() not in special and not inbound.get(f)]
    return {
        "success": True, "pages": len(files),
        "broken_links": broken, "orphans": orphans, "not_in_index": not_indexed,
        "without_sources": no_source, "stale": stale, "open_questions": questions,
        "contradictions": conflicts,
        "inbox_waiting": len(list(_iter_files(root, root / "raw" / "inbox", exts=None))),
        "external_inbox_waiting": len(_external_inbox(brain)),
    }

# ── Token-Middleware ──────────────────────────────────────────────────────────


class TokenAuthMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: Request, call_next):
        if not TOKEN:
            return await call_next(request)
        auth_header = request.headers.get("Authorization", "")
        token_param = request.query_params.get("token", "")
        if auth_header == f"Bearer {TOKEN}" or token_param == TOKEN:
            return await call_next(request)
        return Response("Unauthorized", status_code=401)


# DNS-Rebinding-Schutz aus: Anfragen kommen über LAN-IP / MCP-Proxy mit fremdem Host-Header
app = mcp.streamable_http_app(
    transport_security=TransportSecuritySettings(enable_dns_rebinding_protection=False),
    host="0.0.0.0",
)
app.add_middleware(TokenAuthMiddleware)

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=PORT)
